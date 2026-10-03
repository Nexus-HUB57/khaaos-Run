#!/usr/bin/env python3
"""Revisão editorial conservadora de textos em DOCX do Khaaos Run."""
from __future__ import annotations
import json, re, shutil, time, os
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from openai import OpenAI
from docx import Document

ROOT = Path('/home/ubuntu')
OUT = ROOT / 'saida'
OUT.mkdir(parents=True, exist_ok=True)
FILES = [
    (ROOT/'upload/KhaaosRun-Mundos.docx', OUT/'KhaaosRun_Mundos_Revisado.docx'),
    (ROOT/'upload/KhaaosRun-PoderdeCombate.docx', OUT/'KhaaosRun_Poder_de_Combate_Revisado.docx'),
    (ROOT/'upload/KhaaosRun-Hierarquia.docx', OUT/'KhaaosRun_Hierarquia_Revisada.docx'),
]
MODEL = 'gpt-5-mini'
SYSTEM = """Você é revisor editorial profissional de português brasileiro para documentação de um jogo eletrônico chamado Khaaos Run (KR). Revise ortografia, gramática, pontuação, concordância, clareza, coesão e fluidez. Preserve o sentido, o tom de fantasia militar sombria e o conteúdo integral. Não invente, remova, interprete, corrija nem balanceie regras. Preserve exatamente números, datas, valores, percentuais, níveis, quantidades, siglas, códigos, termos técnicos, nomes próprios, itens, patentes, mundos, ciclos e unidades. Preserve termos do universo, inclusive KR, KRC, XP, Zheers, Onith, Skeetw, Huran, Ven’Ex, Khael, Malakor, OverLord, Marechal, CIGKV, Ckesalth, Arukheea, Priadus, Barburhen, Khaaos Ville, War’I, Ring, Zu, Factho, Forja do Armeiro, Battle Million e demais nomes. Não transforme afirmações do projeto em fatos verificados. Não converta moedas nem altere a notação dos valores. Mantenha a quantidade e a ordem de itens/listas e preserve quebras de linha. Se um trecho estiver ambíguo, mantenha a formulação tão próxima do original quanto possível, corrigindo somente a língua. Retorne todos os IDs recebidos, um para cada texto, e apenas o JSON solicitado."""
SCHEMA = {
  'type':'json_schema','json_schema':{'name':'rewrites','strict':True,'schema':{
    'type':'object','properties':{'items':{'type':'array','items':{'type':'object','properties':{
      'id':{'type':'integer'},'text':{'type':'string'}},'required':['id','text'],'additionalProperties':False}}},
    'required':['items'],'additionalProperties':False}}
}
CODE_RE = re.compile(r'(?m)^\s*(?:class\b|def\b|return\b|import\b|from\b|for\b|if\b|elif\b|else\s*:|try\s*:|except\b|print\s*\(|self\.|@staticmethod|@dataclass|#|"""|[A-Za-z_]\w*\s*=\s*[^=])|(?:self\.|\.append\s*\(|\.items\s*\(|\.get\s*\(|Enum\b|Tuple\[|List\[|Dict\[|__init__|__name__|\+=|-=|\*\*=)')
NUM_RE = re.compile(r"\d+(?:[.,/]\d+)*(?:%|º|°)?")
TERMS = ['Khaaos Run','KRC','XP','Zheers','Onith','Skeetw','Huran','Ven’Ex','Khael','Malakor','OverLord','CIGKV','Ckesalth','Arukheea','Priadus','Barburhen','Khaaos Ville','War’I','Ring','Zu','Factho','Forja do Armeiro','Battle Million','Marechal Henricks Hunterheads III']

def is_code(text: str) -> bool:
    s=text.strip()
    if not s: return True
    if CODE_RE.search(s): return True
    # Preserve code fragments and command/code menus, including paragraphs with several lines.
    if re.search(r"\b(?:class|def|elif|return|self\.|print\(|input\(|\.append\(|\.items\(|\.get\(|lambda|except|import|Enum|dataclass)\b", s): return True
    if re.search(r"^[A-Za-z_][A-Za-z0-9_\.\[\]]*\s*=\s*[^=]", s) and len(s)<220: return True
    return False

def has_visual_or_field(paragraph) -> bool:
    xml=paragraph._p.xml
    return any(x in xml for x in ('<w:drawing','<w:pict','<w:object','w:type="page"','<w:fldChar'))

def numeric_signature(s: str): return Counter(NUM_RE.findall(s))

def check_candidate(src: str, dst: str) -> tuple[bool,str]:
    if numeric_signature(src) != numeric_signature(dst): return False, 'números/valores alterados'
    if src.count('\n') != dst.count('\n'): return False, 'quebras de linha/lista alteradas'
    # Required proper terms that occur in the source must remain in the edited text.
    low=dst.casefold()
    for t in TERMS:
        if t.casefold() in src.casefold() and t.casefold() not in low:
            return False, f'termo protegido removido: {t}'
    if not dst.strip(): return False, 'texto vazio'
    return True,''

def chunk_indices(items, max_chars=4600, max_items=14):
    chunks=[]; chunk=[]; size=0
    for item in items:
        n=len(item['text'])+50
        if chunk and (size+n>max_chars or len(chunk)>=max_items):
            chunks.append(chunk); chunk=[]; size=0
        chunk.append(item); size+=n
    if chunk: chunks.append(chunk)
    return chunks

client=OpenAI()

def revise_chunk(chunk):
    payload=json.dumps(chunk,ensure_ascii=False)
    last=None
    for attempt in range(3):
        try:
            r=client.chat.completions.create(
                model=MODEL,
                messages=[{'role':'system','content':SYSTEM},{'role':'user','content':payload}],
                response_format=SCHEMA,
                max_completion_tokens=6000,
            )
            data=json.loads(r.choices[0].message.content)
            mapped={int(x['id']):x['text'] for x in data['items']}
            expected={int(x['id']) for x in chunk}
            if set(mapped)!=expected: raise ValueError(f'IDs divergentes {set(mapped)^expected}')
            return mapped
        except Exception as e:
            last=e; time.sleep(1.5*(attempt+1))
    raise RuntimeError(f'Falha após tentativas: {last}')

all_logs=[]
for src_path, out_path in FILES:
    print(f'Preparando {src_path.name}', flush=True)
    shutil.copy2(src_path,out_path)
    doc=Document(out_path)
    para_map={}
    eligible=[]
    skipped_code=0; skipped_visual=0
    for i,p in enumerate(doc.paragraphs):
        text=p.text
        if not text.strip(): continue
        if has_visual_or_field(p):
            skipped_visual+=1; continue
        if is_code(text):
            skipped_code+=1; continue
        eligible.append({'id':i,'text':text})
        para_map[i]=p
    chunks=chunk_indices(eligible)
    rewrites={}; failures=[]
    print(f'  {len(eligible)} parágrafos editoriais em {len(chunks)} lotes; {skipped_code} trechos de código e {skipped_visual} elementos visuais/campos preservados.',flush=True)
    with ThreadPoolExecutor(max_workers=5) as pool:
        futures=[pool.submit(revise_chunk,c) for c in chunks]
        for fut in as_completed(futures):
            mapped=fut.result(); rewrites.update(mapped)
            print(f'  lotes concluídos: {len(rewrites)}/{len(eligible)} textos',flush=True)
    accepted=0; reverted=0
    for item in eligible:
        i=item['id']; old=item['text']; new=rewrites.get(i,old).strip()
        ok,reason=check_candidate(old,new)
        if not ok:
            failures.append({'id':i,'reason':reason,'original':old[:300],'proposed':new[:300]})
            new=old; reverted+=1
        elif new!=old:
            accepted+=1
        p=para_map[i]
        runs=p.runs
        if runs:
            runs[0].text=new
            for run in runs[1:]: run.text=''
        else:
            p.add_run(new)
    doc.save(out_path)
    log={'source':str(src_path),'output':str(out_path),'candidate_paragraphs':len(eligible),'rewritten':accepted,'code_paragraphs_preserved':skipped_code,'visual_field_paragraphs_preserved':skipped_visual,'reverted_validation_failures':reverted,'validation_failures':failures,'tables_preserved':len(doc.tables),'images_preserved':len(doc.inline_shapes)}
    log_path=OUT/(out_path.stem+'_review_log.json')
    log_path.write_text(json.dumps(log,ensure_ascii=False,indent=2))
    all_logs.append(log)
    print(f'  salvo: {out_path} | reescritos={accepted}, revertidos por validação={reverted}',flush=True)
(OUT/'kr_revisao_log.json').write_text(json.dumps(all_logs,ensure_ascii=False,indent=2))
print('REVISAO_COMPLETA',flush=True)
